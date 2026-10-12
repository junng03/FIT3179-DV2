"""Build charts 1 and 2 data from the ACMA data tables workbook.

Source: ACMA (2026), Communications and media in Australia: How we watch and listen
to content, data tables (xlsx), question QF8: "Which of the following online services,
if any, have you used to watch video content in the past 7 days?" (adults 18+).

Usage: python scripts/clean_acma.py "path/to/How we watch and listen to content data tables.xlsx"
Writes: data/acma_services_by_year.csv (chart 1), data/acma_ugc_by_age.csv
        (chart 2: QF4 "User-generated or short-form online video service", past 7 days, by age, 2024 vs 2025)

Cleaning rules (nothing is estimated):
- Only the "Column %" row of each service is used (the n, base and significance rows are skipped).
- A value of exactly 0 means the service was not in that year's question list, so it is dropped.
- 2019 is dropped: its question list differed (YouTube 12% vs 44% in 2020; Facebook and TikTok not asked).
- Percentages are rounded to 1 decimal place.
"""
import csv, os, sys
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SOURCE = "ACMA (2026), How we watch and listen to content, data tables, QF8"

# Workbook label -> short name used on the page
SERVICES = {
    "YouTube + YouTube Premium (2020 ONLY)": "YouTube",
    "Netflix": "Netflix",
    "Facebook": "Facebook",
    "Instagram Reels": "Instagram Reels",
    "Amazon Prime Video": "Prime Video",
    "Disney+": "Disney+",
    "TikTok": "TikTok",
    "Stan": "Stan",
}


def pct_rows(ws):
    """Yield (label, values) for each service's Column % row."""
    for row in ws.iter_rows(values_only=True):
        label = row[0]
        if isinstance(label, str) and label.strip() in SERVICES:
            yield SERVICES[label.strip()], row[1:]


def main(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    # Chart 1: by study year
    ws = wb["QF8. Which of by Study year"]
    rows = list(ws.iter_rows(values_only=True))
    years = [c for c in rows[2][1:]]
    out = []
    for name, vals in pct_rows(ws):
        for year, v in zip(years, vals):
            if year is None or v in (None, 0) or int(year) < 2020:
                continue
            out.append({"year": int(year), "service": name, "pct": round(v * 100, 1), "source": SOURCE})
    with open(os.path.join(ROOT, "data", "acma_services_by_year.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["year", "service", "pct", "source"])
        w.writeheader(); w.writerows(out)
    print(f"acma_services_by_year.csv: {len(out)} rows")

    # Chart 2: share watching user-generated or short-form video in the past 7 days, by age, 2024 vs 2025
    ws = wb["QF4. Which of by PBI BANNER"]
    rows = list(ws.iter_rows(values_only=True))
    groups, last = [], None
    for c in rows[3][1:]:
        last = c if c is not None else last
        groups.append(last)
    years = rows[5][1:]
    ages = ["18-24", "25-34", "35-44", "45-54", "55-64", "65-74", "75+"]
    ugc = next(r for r in rows if isinstance(r[0], str) and r[0].startswith("User-generated or short-form"))
    out = [{"age_group": g, "year": int(y), "pct": round(v * 100, 1),
            "source": "ACMA (2026), How we watch and listen to content, data tables, QF4 by age"}
           for g, y, v in zip(groups, years, ugc[1:]) if g in ages and str(y) in ("2024", "2025")]
    with open(os.path.join(ROOT, "data", "acma_ugc_by_age.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["age_group", "year", "pct", "source"])
        w.writeheader(); w.writerows(out)
    print(f"acma_ugc_by_age.csv: {len(out)} rows")


if __name__ == "__main__":
    main(sys.argv[1])
